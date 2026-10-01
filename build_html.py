"""Build the single current, offline HTML presentation from editable slide content."""
import base64
import hashlib
import html
from pathlib import Path

from pptx.oxml.ns import qn

from build_slides import ASSETS, OUT, prs

ROOT = Path(__file__).resolve().parent
PIXELS_PER_EMU = 96 / 914400

assets = {}


def esc(value):
    return html.escape(str(value), quote=True)


def css_color(font):
    try:
        return str(font.color.rgb) if font.color.type else None
    except (AttributeError, TypeError):
        return None


def font_style(font):
    rules = []
    if font.size:
        rules.append(f'font-size:{font.size.pt * 4 / 3:.3f}px')
    if font.name:
        rules.append('font-weight:500' if 'Medium' in font.name else 'font-weight:400')
    if font.bold:
        rules.append('font-weight:600')
    color = css_color(font)
    if color:
        rules.append(f'color:#{color}')
    return ';'.join(rules)


def paragraph_html(paragraph):
    style = font_style(paragraph.font)
    spacing = paragraph.line_spacing
    if isinstance(spacing, float):
        style += f';line-height:{spacing}'
    elif spacing is not None:
        style += f';line-height:{spacing.pt * 4 / 3:.3f}px'
    runs = []
    for run in paragraph.runs:
        run_style = font_style(run.font)
        highlight = run._r.find(qn('a:rPr') + '/' + qn('a:highlight') + '/' + qn('a:srgbClr'))
        if highlight is not None:
            run_style += ';background-color:#' + highlight.get('val')
        words = esc(run.text).replace('\v', '<br>')
        if run.hyperlink.address:
            words = f'<a href="{esc(run.hyperlink.address)}" target="_blank" rel="noopener">{words}</a>'
        runs.append(f'<span style="{esc(run_style)}">{words}</span>')
    return f'<p style="{esc(style)}">{"".join(runs) or "&nbsp;"}</p>'


def picture_asset(shape, slide_number, position):
    if shape.name == 'DragonFly logo':
        blob = (ASSETS / 'brand' / 'dragonfly-mark-soft-black.svg').read_bytes()
        mime = 'image/svg+xml'
    else:
        blob, mime = shape.image.blob, shape.image.content_type
    key = 'asset-' + hashlib.sha256(blob).hexdigest()[:12]
    assets[key] = f'data:{mime};base64,' + base64.b64encode(blob).decode('ascii')
    return key


def shape_html(shape, slide_number, position):
    classes = ['shape']
    attrs = []
    if position == (5 if slide_number == 1 else 4):
        classes.append('heading')
    coords = [f'{name}:{getattr(shape, attr) * PIXELS_PER_EMU:.3f}px'
              for name, attr in [('left', 'left'), ('top', 'top'), ('width', 'width'), ('height', 'height')]]
    style = ';'.join(coords)
    content = ''
    if hasattr(shape, 'image'):
        classes.extend(['graphic-shape', picture_asset(shape, slide_number, position)])
        left, right = shape.crop_left, shape.crop_right
        top, bottom = shape.crop_top, shape.crop_bottom
        if any((left, right, top, bottom)):
            # Render the native picture crop without altering the source image.
            style += (f';background-size:{100/(1-left-right):.4f}% {100/(1-top-bottom):.4f}%'
                      f';background-position:{100*left/(left+right) if left+right else 50:.4f}% '
                      f'{100*top/(top+bottom) if top+bottom else 50:.4f}%')
        if position == 1:
            attrs.append('aria-hidden="true"')
        else:
            labels = {(2, 5): 'Amit Vijapur', (2, 6): 'Jason Cheng',
                      (18, 6): 'Amit LinkedIn QR code', (18, 9): 'Jason LinkedIn QR code'}
            attrs.append('role="img"')
            label = shape.name if shape.name in ('Durham crest', 'OpenAI logo', 'DragonFly logo') else labels.get((slide_number, position), 'Tool logo')
            attrs.append(f'aria-label="{esc(label)}"')
        url = shape.click_action.hyperlink.address
        if url:
            content = f'<a class="image-link" aria-label="Open LinkedIn profile" href="{esc(url)}" target="_blank" rel="noopener"></a>'
    elif shape.has_text_frame and shape.text:
        classes.append('text-shape')
        content = ''.join(paragraph_html(p) for p in shape.text_frame.paragraphs)
    else:
        classes.append('graphic-shape')
        attrs.append('aria-hidden="true"')
        try:
            if shape.fill.type:
                style += ';background-color:#' + str(shape.fill.fore_color.rgb)
        except (AttributeError, TypeError):
            pass
        try:
            if shape.line.fill.type:
                style += ';border:1px solid #' + str(shape.line.color.rgb)
        except (AttributeError, TypeError):
            pass
    return f'<div class="{" ".join(classes)}" style="{esc(style)}" {" ".join(attrs)}>{content}</div>'


def build():
    sections = []
    for number, slide in enumerate(prs.slides, 1):
        title = slide.shapes[4 if number == 1 else 3].text.replace('\n', ' ')
        note = slide.notes_slide.notes_text_frame.text
        pieces = [shape_html(shape, number, pos)
                  for pos, shape in enumerate(slide.shapes, 1)]
        sections.append(f'<section class="slide" aria-label="{number}. {esc(title)}" '
                        f'data-title="{esc(title)}" data-notes="{esc(note)}" '
                        f'aria-hidden="true" inert>'
                        + ''.join(pieces) + '</section>')
    asset_css = '\n'.join(f'.{key}{{background-image:url("{uri}")}}' for key, uri in assets.items())
    template = (ROOT / 'web' / 'shell.html').read_text()
    result = (template.replace('<!-- DECK -->', '\n'.join(sections))
              .replace('/* DECK_CSS */', (ROOT / 'web' / 'deck.css').read_text() + '\n' + asset_css)
              .replace('/* DECK_JS */', (ROOT / 'web' / 'deck.js').read_text()))
    target = OUT / 'vibe-coding-workshop.html'
    target.write_text(result)
    print(f'{target}\n{len(prs.slides)} slides; one automatic fade per slide; offline assets.')


if __name__ == '__main__':
    build()
