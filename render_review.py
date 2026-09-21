from pathlib import Path
import math
import sys
import pypdfium2 as pdfium
from pypdf import PdfReader
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'out'
stem = sys.argv[1] if len(sys.argv) > 1 else 'vibe-coding-workshop-draft-v6'
version = stem.rsplit('-',1)[-1]
prefix = version+'-' if version.startswith('v') and version[1:].isdigit() else ''
pdf=pdfium.PdfDocument(OUT/f'{stem}.pdf')
sheet=Image.new('RGB',(1600,math.ceil(len(pdf)/3)*327),(224,223,216))
draw=ImageDraw.Draw(sheet)
for i,page in enumerate(pdf):
    im=page.render(scale=1).to_pil().convert('RGB')
    im.thumbnail((510,287))
    x=20+(i%3)*530
    y=20+(i//3)*327
    sheet.paste(im,(x,y))
    draw.text((x,y+293),f'{i+1:02}',fill=(30,30,30))
    page.render(scale=1.6).to_pil().save(OUT/f'{prefix}slide-{i+1:02}.png')
sheet.save(OUT/f'{prefix}contact-sheet.jpg',quality=93)
prs=Presentation(OUT/f'{stem}.pptx')
assert len(pdf)==len(prs.slides)
assert all(s.has_notes_slide and len(s.notes_slide.notes_text_frame.text)>40 for s in prs.slides)
reader=PdfReader(OUT/f'{stem}.pdf')
fonts=sorted({str(f.get_object().get('/BaseFont')) for p in reader.pages for f in p['/Resources'].get('/Font',{}).values()})
bounds=[]
for i,s in enumerate(prs.slides):
    for sh in s.shapes:
        if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+10 or sh.top+sh.height>prs.slide_height+10:
            bounds.append((i+1,sh.name))
print(f'PDF pages and PPTX slides: {len(pdf)}. Presenter notes: {len(prs.slides)}/{len(prs.slides)}.')
print('PDF fonts:', fonts)
print('Out-of-slide shapes:',bounds)
print(OUT/f'{prefix}contact-sheet.jpg')
